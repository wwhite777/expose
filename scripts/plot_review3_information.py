#!/usr/bin/env python3
"""Render the Review-3 information-access contract without reading result data."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import subprocess


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


PANEL_CONTENT = (
    ("R1 · historical reference",
     "Training observations\nS source + h labels\nfrom session 1",
     "Same classifier\nhistorical S1 alignment\nfor S + h training",
     "S2 reference: target S1\nall 100 / 144 signals",
     "Score permitted S2 rows\nall 100 / 144 trials", "009E73"),
    ("R2 · current prefix",
     "Training observations\nS source + h labels\nfrom session 1",
     "Same classifier\nhistorical S1 alignment\nfor S + h training",
     "S2 reference: first 20\nunlabeled current trials",
     "Score permitted S2 tail\n80 OpenBMI / 124 BNCI", "CC79A7"),
    ("R2 · current full batch",
     "Training observations\nS source + h labels\nfrom session 1",
     "Same classifier\nhistorical S1 alignment\nfor S + h training",
     "S2 reference: all\n100 / 144 current trials",
     "Score permitted S2 rows\nall 100 / 144\n(transductive)", "CC79A7"),
)
FOOTER = ("Source-person alignment references use their session-1 signals; S is the final source-classifier count.\n"
          "Broader source-only tuning may use additional source data. Session-2 labels score predictions only.")


def write_pptx(path):
    """Create an editable 6.5 × 4.4 in counterpart using only native shapes."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    presentation = Presentation(); presentation.slide_width = Inches(6.5); presentation.slide_height = Inches(4.4)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    width, height, left = Inches(.80), Inches(.54), Inches(.09)
    panel_width = Inches(2.02)
    for index, (title, training, classifier, reference, scoring, color_hex) in enumerate(PANEL_CONTENT):
        x = Inches(.10 + index * 2.13); color = RGBColor.from_string(color_hex)
        heading = slide.shapes.add_textbox(x, Inches(.11), panel_width, Inches(.28))
        paragraph = heading.text_frame.paragraphs[0]; paragraph.text = title; paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(9); paragraph.font.bold = True
        for y, text in ((.46, training), (1.32, classifier), (2.18, reference), (3.04, scoring)):
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(y), panel_width, height)
            box.fill.solid(); box.fill.fore_color.rgb = RGBColor(255, 255, 255)
            box.line.color.rgb = color; box.line.width = Pt(1.0)
            frame = box.text_frame; frame.clear(); frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame.margin_left = frame.margin_right = Inches(.035)
            paragraph = frame.paragraphs[0]; paragraph.text = text; paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.size = Pt(6.5)
        for y in (.98, 1.84, 2.70):
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, x + Inches(.92), Inches(y), Inches(.18), Inches(.22))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = color; arrow.line.color.rgb = color
    footer = slide.shapes.add_textbox(Inches(.12), Inches(3.68), Inches(6.26), Inches(.58))
    frame = footer.text_frame; frame.clear(); frame.word_wrap = True
    for index, line in enumerate(FOOTER.splitlines()):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line; paragraph.alignment = PP_ALIGN.CENTER; paragraph.font.size = Pt(6.5)
    presentation.core_properties.title = "Information access by condition"
    presentation.core_properties.author = ""
    presentation.save(path)


def ensure_pptx(path, fallback_python):
    try:
        write_pptx(path)
    except ModuleNotFoundError as error:
        if error.name != "pptx":
            raise
        completed = subprocess.run([str(fallback_python), str(Path(__file__).resolve()), "--pptx-only", "--out-dir", str(path.parent)],
                                   text=True, capture_output=True, check=False)
        if completed.returncode != 0:
            raise RuntimeError("PPTX fallback failed: " + completed.stdout + completed.stderr)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--pptx-python", type=Path, default=Path("/usr/bin/python3"),
                        help="interpreter used only when the calling interpreter lacks python-pptx")
    parser.add_argument("--pptx-only", action="store_true", help="internal native-PPTX writer; requires an existing output directory")
    args = parser.parse_args()
    if args.pptx_only:
        if not args.out_dir.is_dir() or (args.out_dir / "information_contract.pptx").exists():
            raise FileExistsError("pptx-only requires an existing output directory with no PPTX")
        write_pptx(args.out_dir / "information_contract.pptx")
        return
    if args.out_dir.exists():
        raise FileExistsError("out-dir must be new and exclusive")
    args.out_dir.mkdir(parents=True)

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9,
                         "svg.hashsalt": "review3-information-contract-v1"})
    fig, axes = plt.subplots(1, 3, figsize=(6.5, 4.4))
    green, purple = "#009E73", "#CC79A7"

    def panel(ax, title, training, classifier, reference, scoring, color):
        ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
        ax.set_title(title, fontsize=9, fontweight="bold", pad=7)
        boxes = ((.09, .76, training), (.09, .53, classifier), (.09, .29, reference), (.09, .045, scoring))
        width, height = .82, .15
        for x, y, text in boxes:
            patch = FancyBboxPatch((x, y), width, height, boxstyle="round,pad=0.02,rounding_size=0.025",
                                   facecolor="#ffffff", edgecolor=color, linewidth=1.25)
            ax.add_patch(patch)
            ax.text(.50, y + height / 2, text, ha="center", va="center", fontsize=7.0, linespacing=1.08)
        for _, y, _ in boxes[:-1]:
            ax.annotate("", (.50, y - .055), (.50, y), arrowprops={"arrowstyle":"->", "lw":1.2, "color":color})

    for axis, (title, training, classifier, reference, scoring, color) in zip(axes, PANEL_CONTENT):
        panel(axis, title, training, classifier, reference, scoring, "#" + color)
    fig.text(.5, .006, FOOTER, ha="center", va="bottom", fontsize=6.7, color="#333333", linespacing=1.25)

    caption = ("Information access by condition. The same classifier is trained once using historical session-1 alignment for S plus h. R1 uses a target-session-1 evaluation transform reference and scores all session-2 trials. "
               "R2 prefix uses 20 unlabeled current-session trials for the transform and scores the remaining tail; R2 full batch uses all current-session trials and scores that same batch transductively. "
               "Current references change evaluation whitening only; they do not retrain the classifier. S is the final source fit count, h is the returning person’s session-1 label count, and no outcome values are shown.")
    fig.tight_layout(rect=(0, .055, 1, 1), w_pad=.65)
    fig.savefig(args.out_dir / "information_contract.png", dpi=240, bbox_inches="tight", metadata={"Software":"plot_review3_information"})
    fig.savefig(args.out_dir / "information_contract.pdf", bbox_inches="tight", metadata={"Title":"Information access by condition", "Creator":"plot_review3_information", "CreationDate":None})
    plt.close(fig)
    ensure_pptx(args.out_dir / "information_contract.pptx", args.pptx_python)
    metadata = {"schema":"review3-information-figure-v1", "created_utc":datetime.now(timezone.utc).isoformat(),
                "script_sha256":sha256(__file__), "caption":caption,
                "files":{name:sha256(args.out_dir/name) for name in ("information_contract.png","information_contract.pdf","information_contract.pptx")},
                "content_contract":{"outcomes_read":False,"outcomes_shown":False,"main_grid_current_session_signals":False,
                                    "historical_signal_allowance":{"openbmi":100,"bnci":144},
                                    "prefix_current_session":{"prefix":20,"tail_openbmi":80,"tail_bnci":124},
                                    "full_batch_current_session":{"openbmi":100,"bnci":144,"transductive":True}}}
    (args.out_dir / "FIGURE_INDEX.json").write_text(json.dumps(metadata,indent=2,sort_keys=True)+"\n",encoding="utf-8")

if __name__ == "__main__":
    main()
